#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Read and display Word document structure."""

from docx import Document
import sys

def read_docx(file_path):
    doc = Document(file_path)

    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    print("\n" + "="*80)

    # Show structure with paragraph indices
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            style = para.style.name
            print(f"[{i:3d}] {style:30s} | {text[:100]}")

    print("\n" + "="*80)
    print("\nTables:")
    for i, table in enumerate(doc.tables):
        print(f"\nTable {i+1} ({len(table.rows)} rows x {len(table.columns)} cols):")
        # Show first row
        if table.rows:
            first_row = [cell.text.strip() for cell in table.rows[0].cells]
            print(f"  First row: {first_row[:5]}")

if __name__ == '__main__':
    file_path = r"C:\Users\nguye\Downloads\BAO CAO NHOM 6.docx"
    read_docx(file_path)
