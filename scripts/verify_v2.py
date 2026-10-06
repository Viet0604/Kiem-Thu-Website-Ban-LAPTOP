from docx import Document

PATH = r"C:\Kiểm thử phần mềm\BTL KIEM THU\BAO CAO NHOM 6 (v2).docx"
doc = Document(PATH)

print("=== HEADINGS ===")
for p in doc.paragraphs:
    t = p.text.strip()
    if t and p.style.name.startswith('Heading'):
        print(f"  [{p.style.name}] {t}")

print(f"\n=== TỔNG SỐ BẢNG: {len(doc.tables)} ===")
for i, tbl in enumerate(doc.tables):
    h = tbl.rows[0].cells[0].text.strip()
    print(f"  Bảng {i:02d}: {len(tbl.rows)} rows x {len(tbl.columns)} cols | {h[:50]}")

# Kiểm tra nội dung key
key_checks = [
    'E2E-02', 'E2E-04',
    'IT-AUTH-01', 'IT-INV-01',
    'ST-01', 'ST-08',
    'UT-CART-01', 'UT-CHK-07', 'UT-ORD-02',
    '3.1.', '3.2.', '3.3.', '3.4.',
    'PASS', '100%'
]
body_text = '\n'.join(p.text for p in doc.paragraphs)
for tbl in doc.tables:
    for row in tbl.rows:
        body_text += '\n' + ' '.join(c.text for c in row.cells)

print("\n=== KIỂM TRA NỘI DUNG ===")
for kw in key_checks:
    found = kw in body_text
    status = "✅" if found else "❌"
    print(f"  {status} '{kw}'")
